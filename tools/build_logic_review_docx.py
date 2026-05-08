from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("金森论文复现与硕士论文结构逻辑检查_20260505.docx")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True)
        set_cell_shading(table.rows[0].cells[i], "EAF2F8")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Cm(width)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor(31, 78, 121)


def add_para(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        add_para(doc, item, "List Bullet")


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("金森论文复现与硕士论文结构逻辑检查")
    tr.bold = True
    tr.font.name = "Microsoft YaHei"
    tr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    tr.font.size = Pt(18)
    tr.font.color.rgb = RGBColor(31, 78, 121)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("基于当前论文逻辑稿、结构整理稿与金森 E4-01 论文的中文整理")
    sr.font.name = "Microsoft YaHei"
    sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor(89, 89, 89)

    add_heading(doc, "一、总体判断", 1)
    add_para(
        doc,
        "目前论文主线已经基本成立：以大井町 TOD 站域为对象，将 GPS 停留点作为站域 Place 功能和街道活力的行为代理指标，解释“人在站域中在哪里停留、为什么停留、哪些远离站前的街道仍能产生停留”。这个主线比“老年友好步行环境”更集中，也更容易与金森论文形成可复现、可扩展的关系。",
    )
    add_para(
        doc,
        "但若要达到硕士论文程度，不能只把金森论文的方法换成更多变量。需要把会议论文式的“政策背景-数据-模型-结果”扩展为硕士论文式的“理论定位-先行研究缺口-研究问题-变量体系-模型检验-机制讨论-规划含义”。换句话说，论文的重心应从“我用了 GPS、空间句法、POI、街景”改为“这些变量如何共同解释 TOD 站域从通过空间转化为停留空间”。",
    )

    add_heading(doc, "二、金森论文的可复现骨架", 1)
    add_table(
        doc,
        ["金森论文部分", "实际承担的逻辑功能", "可复现到硕士论文的方式"],
        [
            ["题目与摘要", "明确对象、因变量、核心解释变量和距离视角。因变量是滞在数，解释变量是步道结构、沿道 POI 与距站距离。", "题目保留“空间分布”和“建成环境影响机制”，摘要中明确因变量、分析尺度、变量组和模型。"],
            ["1. はじめに", "从日本 walkable 政策切入，指出站前集聚之外，远离车站区域的滞在创造不足，形成研究 gap。", "Introduction 不能从技术开始，应从 TOD 站域的 Place 功能、公共生活和停留行为切入，再提出远站街道的解释问题。"],
            ["2. 研究手法", "先定义分析单元，再说明步道网络、POI、滞在点、模型。方法服务于问题。", "将研究区、数据、变量、模型分成独立小节，补充阈值选择、数据时间一致性、变量聚合方式和稳健性分析。"],
            ["3. 対象地域", "说明大井町与吉祥寺的交通规模、步道结构、POI 分布和停留分布，为结果解释铺垫。", "大井町单案例也要写清再开发背景、500m/10分钟步行圈、商店街/设施/车站出入口等空间条件。"],
            ["4. 結果", "用 GLM 和评分分析证明距站距离有负效应，但优质 POI 和步行者空间可缓和距离衰减。", "Results 分为描述性空间分布、基准模型、扩展模型、交互项、稳健性检验，避免直接进入解释。"],
            ["5. 考察と結論", "把结果解释为远站区域也能通过空间条件创造滞在，并指出 link 单位分析的价值与限制。", "Discussion 单独成章，解释近站/远站机制差异，说明空间句法和街景变量相对金森的新增贡献。"],
        ],
        [3.0, 6.0, 6.0],
    )

    add_heading(doc, "三、当前论文逻辑检查", 1)
    add_bullets(
        doc,
        [
            "主线优点：当前稿已把研究定位从一般 walkability 收束为 TOD 站域内的 staying behaviour，并且把 GPS 停留点定义为 Place 功能的行为代理指标，这是可以支撑硕士论文的核心概念。",
            "主要风险一：摘要和目录中工具名称较多，理论问题还需要前置。应先解释为什么停留行为比通过量更接近街道场所品质，再进入 GPS、空间句法、POI 和街景。",
            "主要风险二：金森使用两个站点进行比较，而你目前更像大井町单案例。单案例可以成立，但必须加强研究区选择理由、再开发背景、边界设定和大井町内部空间差异。",
            "主要风险三：如果模型尚未完成，摘要中的“结果表明”应谨慎；可以写为“旨在检验”“预期揭示”。模型完成后再改成实证结论。",
            "主要风险四：Results 与 Discussion 必须分开。Results 报告分布、模型和稳健性；Discussion 再解释距离衰减、远站停留、设施质量、街景体验和规划含义。",
        ],
    )

    add_heading(doc, "四、硕士论文建议目录", 1)
    add_table(
        doc,
        ["章节", "建议标题", "核心问题", "应放入的材料"],
        [
            ["第1章", "绪论", "为什么 TOD 站域需要研究停留，而不只是研究通行和可达性？", "研究背景、问题意识、研究目的、研究问题、贡献、论文结构。"],
            ["第2章", "先行研究与理论框架", "既有研究如何理解 Node-Place、Link-Place、walkability 与停留行为？缺口在哪里？", "TOD/Node-Place、Link-Place、公共生活、停留行为、GPS/POI/街景/空间句法相关研究。"],
            ["第3章", "研究对象与数据", "为什么选择大井町，数据如何被统一到 street segment？", "研究区边界、再开发背景、GPS、POI、街景、路网/空间句法数据，时间范围和预处理。"],
            ["第4章", "研究方法与变量设计", "如何从停留点构造因变量，如何构造建成环境解释变量？", "停留点识别算法、segment 匹配、POI 缓冲、街景语义指标、空间句法指标、模型设定。"],
            ["第5章", "结果", "停留在空间上如何分布？哪些变量显著影响停留强度？", "核密度图、segment 停留强度图、圈层统计、负二项回归、交互项、稳健性。"],
            ["第6章", "讨论", "为什么某些远站街道仍能产生停留？本研究相对金森扩展了什么？", "近站/远站机制、空间结构-功能配置-微观体验三层解释、规划启示、与金森对照。"],
            ["第7章", "结论", "研究回答了什么，还有什么限制？", "主要发现、理论贡献、方法贡献、规划建议、局限与未来研究。"],
        ],
        [1.6, 3.6, 5.2, 5.6],
    )

    add_heading(doc, "五、现有内容可以填到哪一部分", 1)
    add_table(
        doc,
        ["现有内容", "可放入章节", "建议处理方式"],
        [
            ["标题、中文摘要、英文摘要、关键词", "论文前置部分；第1章末尾也可复用研究目的与贡献表述", "保留主线，但摘要中暂时避免过强结果表述；模型完成后再更新为实证发现。"],
            ["“当前结构诊断”", "不直接进入正文；作为写作备忘或导师沟通材料", "其中“问题驱动而非技术堆叠”的判断可转化为第1章研究目的和第2章研究缺口。"],
            ["“建议标题与研究定位”", "第1章 研究目的、研究问题、研究贡献", "可直接改写为研究定位段落，尤其是“把经过的人转化为停留的人”这一问题意识。"],
            ["“研究问题与假设”", "第1章 研究问题；第4章 模型假设", "将 RQ 与 H 分开：RQ 放绪论，H 放方法章模型设定前。"],
            ["“修改后的论文目录”", "全文目录基础", "建议按七章结构重排，并将 Results / Discussion 拆开。"],
            ["Introduction 段落骨架", "第1章 绪论", "可以直接扩写为 3-4 个小节：背景、问题、目的、贡献。"],
            ["Literature Review 写作要点", "第2章 先行研究", "按理论到方法排序：TOD/Node-Place -> Link-Place/公共生活 -> staying -> 方法研究。"],
            ["Study Area and Data 要点", "第3章 研究对象与数据", "补充大井町选择理由、边界图、数据时间范围、数据统一到 segment 的流程图。"],
            ["Methodology 要点", "第4章 研究方法与变量设计", "保留负二项回归建议，新增变量定义表、公式、阈值依据、稳健性设计。"],
            ["变量与分析单元设计", "第4章 核心内容", "整理为一张变量表：因变量、距离变量、空间句法、POI、街景、控制变量、预期方向。"],
            ["模型实施建议", "第4章末尾与第5章开头", "分层模型顺序很好，可作为实证分析路线：基准模型、变量组模型、交互模型、稳健性模型。"],
            ["先行研究对应表", "第2章末尾或附录", "正文中不宜只放表，应把表转化为文献综述叙述，并保留表作为写作检查表。"],
            ["可直接改写进正文的段落骨架", "第1章与第2章之间的过渡段", "可以使用，但要补引用，并避免与摘要重复。"],
            ["下一步写作与数据检查清单", "不放正文；作为研究管理清单", "用于检查数据年份、GPS 阈值、POI 缓冲、街景采样点和模型诊断。"],
        ],
        [4.2, 4.0, 6.8],
    )

    add_heading(doc, "六、变量与模型应扩展到硕士论文程度", 1)
    add_table(
        doc,
        ["模块", "金森论文做法", "硕士论文建议扩展"],
        [
            ["因变量", "步道 link 单位的 7 日滞在数，offset 校正步道长度。", "保留停留数，并补充单位长度停留密度、时间段停留、平日/周末差异作为辅助分析。"],
            ["距离变量", "车站重心到 link 中点的经路最短距离，并与变量做交互。", "进一步说明采用车站重心、出入口或站前广场作为起点的理由；可做距离定义敏感性。"],
            ["POI", "总 POI、热门 POI、高评价 POI。", "区分数量、质量和类型；可以合成商业/餐饮/公共设施/休憩设施变量，避免 POI 总量解释过粗。"],
            ["步道结构", "步行者专用道路、步车分离、路侧带、步车一体。", "若没有同等数据，可用道路等级、人行空间宽度、交叉口密度、街道断面特征替代并说明限制。"],
            ["空间句法", "金森未纳入。", "作为你的扩展重点，加入 integration、choice、connectivity 等，解释网络到达潜力与穿行潜力。"],
            ["街景环境", "金森未纳入。", "作为第二个扩展重点，加入 GVI、天空开敞度、建筑围合、店面/绿化/道路可视比例等微观体验变量。"],
            ["模型", "负二项 GLM + 评分分析。", "保留负二项作为主模型，补充分层模型、交互项、VIF、Moran's I、阈值敏感性；必要时补空间误差模型或 GWR。"],
        ],
        [3.0, 5.4, 6.6],
    )

    add_heading(doc, "七、建议马上补写的正文块", 1)
    add_bullets(
        doc,
        [
            "第1章补写：从日本 walkable 政策和 TOD 站域公共生活切入，说明停留行为是 Place 功能的行为表现。",
            "第2章补写：把文献综述改成“理论-行为-方法”三段，而不是按工具罗列。",
            "第3章补写：大井町为何适合作为单案例，研究边界如何确定，数据年份是否一致。",
            "第4章补写：停留点识别阈值、segment 匹配规则、变量表、模型公式和假设方向。",
            "第5章预留：先做描述性图，再做模型表，最后做交互项图，不要一开始就讨论意义。",
            "第6章预留：围绕“距离衰减能否被空间结构、功能配置、微观体验削弱”组织讨论。",
        ],
    )

    add_heading(doc, "八、可直接采用的研究问题版本", 1)
    add_para(doc, "RQ1：大井町 TOD 站域内的停留行为在 street segment 尺度上呈现怎样的空间分布和距离衰减特征？")
    add_para(doc, "RQ2：空间句法所刻画的路网结构、POI 所刻画的功能配置、街景语义所刻画的微观体验，分别如何影响停留强度？")
    add_para(doc, "RQ3：在远离站前核心区的街道中，哪些建成环境条件能够削弱距站距离对停留行为的负向影响？")
    add_para(doc, "RQ4：相较于只关注步道结构和沿道设施的金森研究，整合路网拓扑与街景微观环境是否能更充分解释 TOD 站域的停留机制？")

    add_heading(doc, "九、结论性建议", 1)
    add_para(
        doc,
        "当前内容已经可以进入硕士论文写作，但应先把“论文目录”和“现有材料归位”完成，再继续扩写正文。最优先的动作不是增加更多先行研究，而是固定七章结构、确定研究边界、完成变量表和模型路线。只要保持“停留行为作为 TOD Place 功能的行为代理指标”这一主轴，金森论文就可以作为方法和问题意识的参照，而你的硕士论文则通过空间句法与街景微观环境实现扩展。",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
