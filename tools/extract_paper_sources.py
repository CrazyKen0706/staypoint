from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


DOCX_FILES = [
    Path(r"D:/修论/Paper/2021.04.21_标题摘要目录逻辑修改_20260505.docx"),
    Path(r"D:/修论/Paper/硕士论文结构整理与先行研究补强_20260505.docx"),
]
PDF_FILE = Path(r"C:/Users/14572/Zotero/storage/3DRZ48E5/E4-01金森.pdf")
OUT_DIR = Path("analysis_outputs")


def compact(text: str) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    return text


def docx_outline(path: Path) -> dict:
    doc = Document(str(path))
    paragraphs = []
    headings = []
    for i, para in enumerate(doc.paragraphs, 1):
        text = compact(para.text)
        if not text:
            continue
        style = para.style.name if para.style is not None else ""
        item = {"index": i, "style": style, "text": text}
        paragraphs.append(item)
        if "Heading" in style or re.match(r"^第[一二三四五六七八九十0-9]+[章節节]|^[0-9]+(\.[0-9]+)*\s+", text):
            headings.append(item)

    tables = []
    for ti, table in enumerate(doc.tables, 1):
        rows = []
        for row in table.rows[:12]:
            rows.append([compact(cell.text) for cell in row.cells])
        tables.append({"index": ti, "rows": rows})

    return {
        "file": str(path),
        "paragraph_count": len(paragraphs),
        "headings": headings,
        "paragraphs": paragraphs,
        "tables": tables,
    }


def pdf_extract(path: Path) -> dict:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        pages.append({"page": i, "text": compact(text)})
    full_text = "\n".join(p["text"] for p in pages)
    lines = [compact(x) for x in re.split(r"[\n\r]+", full_text) if compact(x)]
    candidates = []
    for line in lines:
        if re.search(r"^(Abstract|Keywords|[0-9]+\.|第[一二三四五六七八九十0-9]+|はじめに|おわりに|まとめ|考察|方法|結果|結論)", line):
            candidates.append(line)
        elif re.search(r"(目的|方法|結果|考察|結論|研究|調査|分析|対象)", line) and len(line) < 80:
            candidates.append(line)
    return {
        "file": str(path),
        "page_count": len(reader.pages),
        "pages": pages,
        "candidate_headings": candidates[:300],
    }


def main() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    data = {"docx": [docx_outline(path) for path in DOCX_FILES], "pdf": pdf_extract(PDF_FILE)}
    (OUT_DIR / "sources_extracted.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    summary_lines = []
    for doc in data["docx"]:
        summary_lines.append(f"# {Path(doc['file']).name}")
        summary_lines.append(f"paragraphs: {doc['paragraph_count']}")
        summary_lines.append("## headings")
        for h in doc["headings"]:
            summary_lines.append(f"- [{h['index']}] {h['style']}: {h['text']}")
        summary_lines.append("## first paragraphs")
        for p in doc["paragraphs"][:80]:
            summary_lines.append(f"- [{p['index']}] {p['style']}: {p['text']}")
        summary_lines.append("")

    pdf = data["pdf"]
    summary_lines.append(f"# {Path(pdf['file']).name}")
    summary_lines.append(f"pages: {pdf['page_count']}")
    summary_lines.append("## candidate headings")
    for h in pdf["candidate_headings"][:120]:
        summary_lines.append(f"- {h}")
    summary_lines.append("## page previews")
    for p in pdf["pages"][:12]:
        summary_lines.append(f"### page {p['page']}")
        summary_lines.append(p["text"][:2500])
    (OUT_DIR / "sources_summary.md").write_text("\n".join(summary_lines), encoding="utf-8")


if __name__ == "__main__":
    main()
