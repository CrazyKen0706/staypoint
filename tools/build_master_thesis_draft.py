from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("大井町TOD站域停留行为研究_硕士论文初稿_20260505.docx")


def font_run(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_para(doc, text="", style=None, size=10.5, bold=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(6)
    if text:
        r = p.add_run(text)
        font_run(r, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        font_run(run, size=15 if level == 1 else 12.5, bold=True, color=(31, 78, 121))
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        font_run(r, size=9, bold=True)
        shade(c, "EAF2F8")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(text)
            font_run(r, size=9)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return table


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.4)
    styles = doc.styles
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    return doc


def build():
    doc = setup_doc()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("大井町 TOD 站域内停留行为的空间分布及建成环境影响机制研究")
    font_run(r, size=18, bold=True, color=(31, 78, 121))
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("硕士论文初稿")
    font_run(r, size=12, color=(89, 89, 89))
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("初稿日期：2026年5月5日")
    font_run(r, size=10, color=(89, 89, 89))

    add_heading(doc, "摘要", 1)
    add_para(doc, "TOD 站域不仅承担公共交通换乘和步行通行功能，也承载等待、休憩、消费、社交与偶遇等日常公共生活活动。既有 TOD 与 walkability 研究多从密度、功能混合、交通可达性、步行连通性等角度评价站域空间，但对于“人在站域街道中是否停留、在哪里停留、为什么停留”的行为层面解释仍相对不足。特别是在东京这样以车站为中心高度集约的都市空间中，停留活动是否过度集中于站前核心区，以及距离车站较远的街道在何种条件下仍能产生停留，是理解站域 Place 功能和街道活力的重要问题。")
    add_para(doc, "本研究以东京都大井町站周边地区为对象，将 GPS 人流数据识别出的停留点作为站域 Place 功能的行为代理指标，并在 street segment 尺度上整合距站距离、空间句法、POI 功能配置、街景微观环境等变量，分析 TOD 站域内停留行为的空间分布及其建成环境影响机制。研究首先识别不同距离圈层和时间段中的停留活动分布特征，其次构建线段尺度的建成环境变量体系，并通过计数回归模型检验空间结构、商业设施与街景环境对停留强度的影响。")
    add_para(doc, "本研究预期揭示：距车站距离通常对停留行为具有负向影响，但在路网整合度较高、沿道设施较丰富、街景环境质量较好的街道中，距离衰减效应可能被削弱，远离站前核心区的街道仍有可能形成较高停留强度。本文旨在从“空间结构-功能配置-微观体验”三个层面解释 TOD 站域内停留行为的形成机制，并为大井町站周边再开发与步行环境优化提供依据。")
    add_para(doc, "关键词：TOD站域；停留行为；空间句法；POI；街景语义分割；GPS人流数据；大井町", bold=True)

    add_heading(doc, "Abstract", 1)
    add_para(doc, "Transit-oriented development (TOD) station areas are not only transport nodes for transfer and pedestrian movement, but also urban places where waiting, resting, consumption, and social encounters occur. Existing TOD and walkability studies have mainly evaluated station areas through density, land-use mix, accessibility, and connectivity, while less attention has been paid to where people actually stay and what built environmental conditions support staying behaviour at the street level.")
    add_para(doc, "Focusing on the area surrounding Oimachi Station in Tokyo, this study regards GPS-derived stay points as behavioural indicators of the place function and street vitality of a TOD station area. At the street-segment level, the study integrates distance from the station, space syntax indicators, POI-based functional variables, and street-view-based micro-environmental indicators to examine the spatial distribution and built environmental determinants of pedestrian staying behaviour.")
    add_para(doc, "The study aims to examine whether streets with higher network integration, richer roadside facilities, and better micro-scale environmental quality can weaken the distance-decay effect and support staying activities outside the immediate station-front area. It contributes to understanding the formation mechanism of staying behaviour in TOD station areas and provides planning implications for walkable station-area regeneration.")
    add_para(doc, "Keywords: TOD station area; staying behaviour; space syntax; POI; street view imagery; GPS human flow data; Oimachi", bold=True)

    add_heading(doc, "第1章 绪论", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_para(doc, "以公共交通节点为中心进行高密度、复合化开发，是现代都市空间组织的重要方式。TOD 强调土地利用与公共交通的协同，通过提高站域密度、功能混合度和步行可达性，促进公共交通使用并减少对汽车的依赖。在东京等轨道交通高度发达的城市中，车站不仅是交通节点，也是商业、服务、公共活动和日常交往高度集聚的城市场所。")
    add_para(doc, "然而，站域空间的价值并不只体现在人流通过和换乘效率上。对于使用者而言，车站周边街道也是等待、休憩、购物、餐饮、社交和偶遇发生的公共生活空间。换言之，TOD 站域既具有 Node 功能，也具有 Place 功能。若研究仅关注步行可达性或通行量，可能难以解释街道是否真正具有吸引人停留的场所品质。")
    add_para(doc, "近年来，日本推进 walkable city 与居心地が良く歩きたくなるまちなか建设，政策重点逐渐从单纯改善通行环境转向创造可停留、可活动、可消费和可交流的城市空间。在这一背景下，如何识别能够支撑停留行为的街道条件，成为站域更新和步行环境优化的重要议题。")

    add_heading(doc, "1.2 研究问题", 2)
    add_para(doc, "既有研究已经较充分地讨论了 TOD 站域的密度、功能混合、公共交通可达性与步行连通性，但对于停留行为的解释仍存在不足。一方面，停留行为不同于移动行为，它更直接反映行人是否愿意在某一街道环境中花费时间；另一方面，站域停留并不必然只发生在站前核心区，部分远离车站的街道也可能因良好的空间结构、沿道功能或街景环境而产生持续活动。")
    add_para(doc, "金森等的研究以大井町站和吉祥寺站周边为对象，使用步道 link 单位的人流停留数，检验了距站距离、步道结构和沿道 POI 对停留的影响。该研究表明，距站距离对停留具有显著负向影响，但优质 POI 和步行者空间条件能够在一定程度上支撑远站区域的停留。该研究为本论文提供了重要方法参照，但仍可进一步扩展：其一，路网拓扑结构如何影响停留潜力尚未被充分纳入；其二，行人在眼平视角中感知到的绿化、开敞度、围合度、街道界面等微观环境仍有待解释；其三，大井町作为再开发背景下的单案例，其内部街道差异值得更细致分析。")

    add_heading(doc, "1.3 研究目的与研究问题", 2)
    add_para(doc, "本研究的目的，是以大井町 TOD 站域为对象，在 street segment 尺度上整合 GPS 停留点、空间句法、POI 与街景环境指标，分析站域内停留行为的空间分布及其建成环境影响机制。论文关注的不是“能否用多种数据描述站域”，而是“站域空间中哪些位置能够把经过的人转化为停留的人”。")
    add_table(doc, ["编号", "研究问题"], [
        ["RQ1", "大井町 TOD 站域内的停留行为在 street segment 尺度上呈现怎样的空间分布和距离衰减特征？"],
        ["RQ2", "空间句法所刻画的路网结构、POI 所刻画的功能配置、街景语义所刻画的微观体验，分别如何影响停留强度？"],
        ["RQ3", "在远离站前核心区的街道中，哪些建成环境条件能够削弱距站距离对停留行为的负向影响？"],
        ["RQ4", "相较于只关注步道结构和沿道设施的金森研究，整合路网拓扑与街景微观环境是否能更充分解释 TOD 站域的停留机制？"],
    ], [2.0, 13.0])

    add_heading(doc, "1.4 研究贡献", 2)
    add_para(doc, "本研究的贡献主要包括三个方面。第一，在行为测量上，将 GPS 数据识别出的停留点作为站域 Place 功能的行为代理指标，使 TOD 研究从通行和可达性评价进一步转向公共生活和空间活力评价。第二，在变量整合上，在线段尺度上同时纳入距站距离、空间句法、POI 与街景环境指标，从空间结构、功能配置和微观体验三个层面解释停留行为。第三，在规划意义上，研究试图识别远离站前核心区但仍能产生停留的街道条件，为大井町站周边再开发、步行环境优化和街道活力提升提供更具体的空间依据。")

    add_heading(doc, "第2章 先行研究与理论框架", 1)
    add_heading(doc, "2.1 TOD、Node-Place 与站域场所功能", 2)
    add_para(doc, "TOD 研究通常强调交通节点与土地利用的协同关系。Bertolini 提出的 Node-Place 模型将车站地区理解为交通节点功能与城市场所功能的复合体，强调站域评价不能只看交通可达性，也要考察其作为城市活动承载地的能力。后续研究将该框架用于不同城市和站域类型，讨论节点功能与场所功能之间的平衡关系。")
    add_para(doc, "对于本研究而言，Node-Place 模型的重要启示在于：站域 Place 功能不能只用土地利用、商业面积或设施数量等静态指标来代理。真正的 Place 功能还应体现在人是否愿意在站域街道中停留、使用和互动。因此，停留行为可以被视为站域 Place 功能的行为化表达。")
    add_heading(doc, "2.2 Link-Place、公共生活与停留行为", 2)
    add_para(doc, "街道既是 movement channel，也是 urban place。Link-Place 框架指出，街道规划不能只强调交通通行功能，还需要兼顾街道作为场所的社会、商业和公共生活功能。与通行行为相比，停留行为更接近街道的场所品质，因为停留意味着使用者愿意在某一空间中花费时间，并与周边设施、界面和环境发生更密切的关系。")
    add_para(doc, "公共生活研究同样强调，街道活力不仅取决于人流量，也取决于停留、观看、交谈、等待和消费等活动能否发生。对于 TOD 站域而言，高人流并不必然等于高活力；如果人只是快速通过，站域仍可能缺乏公共生活。因此，本研究将停留行为作为解释站域街道活力的重要切入点。")
    add_heading(doc, "2.3 停留行为的影响因素研究", 2)
    add_para(doc, "既有关于停留或滞留行为的研究主要关注商业街、公共空间、大规模开发地区和步行环境等对象。相关研究表明，设施吸引力、店铺界面、休憩空间、步道结构、交通安全、绿化和开放空间等因素均可能影响停留活动。金森等以大井町和吉祥寺站周边为对象，使用步道 link 单位的人流停留数，证明距站距离具有稳定负向影响，但优质 POI 和步行者空间条件能够在远站区域支撑一定停留。")
    add_para(doc, "不过，现有研究仍有两个可扩展方向。第一，街道在整体路网中的位置可能影响其可达潜力与经过潜力，进而影响停留机会。第二，街景环境所提供的视觉舒适性、绿化、开敞度和界面连续性，也可能影响人是否愿意停留。基于此，本研究在金森研究的基础上进一步引入空间句法和街景语义指标。")
    add_heading(doc, "2.4 GPS、POI、空间句法与街景数据在城市研究中的应用", 2)
    add_para(doc, "大规模 GPS 数据可以在较细空间尺度上捕捉人的移动和停留行为。与传统问卷或观察相比，GPS 数据具有覆盖范围广、时间连续性强和可与空间数据叠加的优势，但也需要谨慎处理采样间隔、停留阈值、隐私保护和匹配误差等问题。")
    add_para(doc, "POI 数据常用于刻画城市功能配置和设施吸引力，空间句法则用于分析街道网络的整合度、选择度和连接关系，街景图像及语义分割方法能够从行人视角刻画绿化、天空、建筑、道路、店面等微观环境。将这些数据统一到 street segment 尺度，有助于把行为结果与建成环境条件进行更直接的关联分析。")
    add_heading(doc, "2.5 理论框架", 2)
    add_para(doc, "本文构建“空间结构-功能配置-微观体验”的解释框架。空间结构指 street segment 在整体路网中的拓扑位置，影响人流到达潜力和穿行潜力；功能配置指沿道 POI 的数量、类型和质量，影响停留目的和活动机会；微观体验指行人在眼平视角中感知到的街景环境，影响停留的舒适性和吸引力。三者共同作用于停留行为，并可能调节距站距离带来的衰减效应。")

    add_heading(doc, "第3章 研究对象与数据", 1)
    add_heading(doc, "3.1 研究区概况", 2)
    add_para(doc, "本研究以东京都品川区大井町站周边地区为研究对象。大井町站由 JR、东急电铁和东京临海高速铁道等线路服务，是东京南部重要的交通节点。车站周边集聚商业设施、办公、居住、公共服务和再开发项目，同时存在商店街、住宅街、车站前广场、主干道和细街路等多种街道空间类型。")
    add_para(doc, "选择大井町作为研究对象，主要基于三个理由。第一，大井町具有典型的 TOD 站域特征，公共交通可达性高，功能混合明显。第二，车站周边正在经历更新和再开发，识别停留行为与街道环境关系具有现实规划意义。第三，金森研究已经将大井町作为案例之一，本研究可以在其基础上进行方法复现和变量扩展。")
    add_heading(doc, "3.2 研究边界与分析单元", 2)
    add_para(doc, "研究边界建议以大井町站为中心的 500m 圈或 10 分钟步行圈为基础，并结合实际再开发范围进行调整。若 500m 圈与再开发范围不完全一致，正文中应明确主分析边界和补充分析边界。所有数据最终统一到 street segment 或 sidewalk link 单位，以便将停留点、POI、街景采样点和空间句法指标进行一致聚合。")
    add_heading(doc, "3.3 数据来源与处理", 2)
    add_table(doc, ["数据类型", "主要内容", "处理方式", "正文中需补充的信息"], [
        ["GPS 人流数据", "用户位置轨迹与停留点", "依据时间与空间阈值识别停留点，并匹配到最近或可达 segment", "数据年份、时间范围、采样间隔、隐私处理、停留阈值依据"],
        ["道路/步行网络", "街道线段、交叉口、步道结构", "清理拓扑后作为 segment 分析单元", "数据来源、线段分割规则、是否包含站内通路/公园通路"],
        ["POI 数据", "餐饮、商业、服务、文化、休憩等设施", "按 segment 缓冲区聚合，计算数量、密度、类型或质量指标", "POI 取得日期、分类规则、缓冲距离"],
        ["街景数据", "街道眼平视角图像", "语义分割后聚合 GVI、天空、建筑、道路、界面等比例", "街景年份、采样间隔、图像方向、语义分割模型"],
        ["空间句法数据", "integration、choice、connectivity 等", "基于路网计算并赋值到 segment", "半径设定、计算软件、标准化方法"],
    ], [2.5, 3.0, 4.8, 4.7])
    add_para(doc, "需要特别注意的是，不同数据的时间点可能并不完全一致。若 GPS、POI 和街景图像的年份不同，应在正文中说明差异，并在局限性中讨论其可能影响。")

    add_heading(doc, "第4章 研究方法与变量设计", 1)
    add_heading(doc, "4.1 停留点识别与因变量构造", 2)
    add_para(doc, "本研究将 GPS 数据中满足一定空间稳定性和时间持续性的轨迹片段识别为停留点。具体阈值应根据数据采样间隔和定位精度设定，例如在一定距离范围内连续出现若干点，并持续超过最低时间阈值时判定为停留。识别后的停留点通过最近距离或网络可达规则匹配到 street segment，并按 segment 汇总停留数。")
    add_para(doc, "主因变量建议设定为 segment 单位的停留数。考虑不同线段长度可能导致可承载停留机会不同，模型中可使用 log(segment 长度) 作为 offset 项。作为稳健性或补充分析，也可构造单位长度停留密度、平日/周末停留数、白天/夜间停留数等指标。")
    add_heading(doc, "4.2 解释变量设计", 2)
    add_table(doc, ["变量组", "变量示例", "理论含义", "预期方向"], [
        ["距离变量", "距车站经路距离、圈层变量", "反映站前核心区到远站区域的距离衰减", "负向"],
        ["空间句法", "integration、choice、connectivity", "反映街道在路网中的到达潜力和穿行潜力", "多为正向"],
        ["POI 功能配置", "POI 密度、餐饮/商业/休憩设施、高评价 POI", "反映活动目的、设施吸引力和停留机会", "正向，但需考虑类型差异"],
        ["街景微观环境", "GVI、天空开敞度、围合度、店面界面、道路可视比例", "反映行人眼平视角的舒适性和空间体验", "方向视指标而定"],
        ["控制变量", "segment 长度、道路等级、交叉口密度等", "控制线段规模和基础空间差异", "视变量而定"],
        ["交互项", "距站距离 × integration；距站距离 × POI质量；integration × GVI", "检验远站区域中距离衰减是否被环境条件削弱", "重点关注正向调节"],
    ], [2.5, 3.8, 5.2, 3.0])
    add_heading(doc, "4.3 模型设定", 2)
    add_para(doc, "由于 segment 单位停留数属于计数数据，并可能存在过度离散，本研究建议以负二项回归作为主模型。基本模型如下：")
    add_para(doc, "Stay_i ~ NB(μ_i, θ)，log(μ_i) = β0 + β1Distance_i + β2Syntax_i + β3POI_i + β4StreetView_i + β5Controls_i + log(Length_i)")
    add_para(doc, "其中 Stay_i 表示第 i 个 segment 的停留数，Distance_i 表示距站距离，Syntax_i 表示空间句法指标，POI_i 表示沿道设施变量，StreetView_i 表示街景微观环境变量，Controls_i 表示控制变量，log(Length_i) 为 offset 项。")
    add_para(doc, "实证分析可采用分层建模策略：模型1仅加入距站距离和控制变量；模型2加入空间句法变量；模型3加入 POI 变量；模型4加入街景变量；模型5加入关键交互项。通过比较模型拟合度、系数方向和显著性，判断不同变量组对停留行为解释力的增量贡献。")
    add_heading(doc, "4.4 稳健性分析", 2)
    add_para(doc, "为避免结果受到阈值和空间匹配规则影响，本研究应进行稳健性检验。具体包括：改变停留识别阈值、改变 POI 缓冲距离、改变距站圈层划分方式、检查是否剔除站前异常高值后结果仍保持一致。此外，还应检查 VIF 以判断多重共线性，并对模型残差进行 Moran's I 检验，以确认是否存在明显空间自相关。")

    add_heading(doc, "第5章 结果", 1)
    add_heading(doc, "5.1 停留行为的空间分布", 2)
    add_para(doc, "本节应首先展示 GPS 停留点在大井町站域内的总体空间分布。建议使用停留点核密度图、segment 停留强度图和距站圈层统计图，说明停留是否集中于站前核心区，以及远离车站的哪些街道仍出现较高停留。")
    add_para(doc, "待填结果：根据实际图表，描述高停留 segment 的位置、是否沿商店街或主要步行路径分布、是否与再开发地区或大型设施相邻，以及远站区域是否存在局部停留热点。")
    add_heading(doc, "5.2 距站距离与停留强度", 2)
    add_para(doc, "金森研究表明，距站距离通常是解释停留数的重要负向因素。本研究应先检验大井町站域中是否存在类似距离衰减。可按 100m 或更适合研究区的距离圈层计算平均停留数、单位长度停留密度和停留分布离散程度。")
    add_para(doc, "待填结果：若距离衰减明显，应说明衰减速度和转折距离；若远站区域仍存在停留，应指出这些 segment 的空间条件，并为后续交互项模型铺垫。")
    add_heading(doc, "5.3 建成环境变量的描述统计", 2)
    add_para(doc, "本节应报告主要解释变量的描述统计，包括空间句法指标、POI 指标和街景指标的均值、标准差、最小值和最大值。还应通过相关系数矩阵或散点图检查变量之间关系，避免高度相关变量同时进入模型造成解释不稳定。")
    add_heading(doc, "5.4 回归模型结果", 2)
    add_para(doc, "回归结果应按照分层模型顺序呈现。首先报告距站距离的基准效应，其次说明空间句法、POI 和街景变量加入后模型解释力是否提高，再重点解释显著变量的方向和含义。")
    add_para(doc, "待填结果：如果 integration 或 choice 显著为正，可解释为路网中更易到达或更可能经过的街道具有更高停留机会；如果某类 POI 显著为正，可解释为设施吸引力创造停留目的；如果 GVI、开敞度或街道界面变量显著，则可说明微观体验影响停留意愿。")
    add_heading(doc, "5.5 交互项与远站停留机制", 2)
    add_para(doc, "本研究的重点不只是证明近站停留多，而是识别距离衰减能否被某些空间条件削弱。因此，应重点报告距站距离与空间句法、POI 质量、街景质量之间的交互项。若交互项显著，建议绘制边际效应图，展示不同环境条件下距站距离对停留数的影响差异。")
    add_para(doc, "待填结果：若“距站距离 × 高质量 POI”或“距站距离 × integration”为正，说明远站区域中良好的设施或网络位置能够缓和距离衰减；若交互项不显著，则应讨论大井町是否比吉祥寺更依赖站前集聚。")

    add_heading(doc, "第6章 讨论", 1)
    add_heading(doc, "6.1 距离衰减与站域停留的空间集中", 2)
    add_para(doc, "若模型结果支持距站距离的负向效应，可以说明大井町站域的停留活动仍然受到车站核心区吸引力的强烈影响。这一结果与金森研究相一致，表明即使在功能混合度较高的 TOD 站域中，车站仍然是停留活动的主要中心。")
    add_para(doc, "但距离衰减并不意味着远站区域没有规划价值。相反，远站区域中仍能产生停留的街道，往往更能反映建成环境条件本身的作用。站前区域的停留可能由交通节点、大型设施和高人流共同驱动，而远站街道的停留则更依赖空间结构、沿道功能和微观体验的综合支撑。")
    add_heading(doc, "6.2 空间结构、功能配置与微观体验的综合作用", 2)
    add_para(doc, "本研究的理论框架强调，停留行为并非由单一因素决定。空间句法指标反映街道在路网中的位置，影响人流到达和经过机会；POI 指标反映沿道活动目的和设施吸引力；街景指标反映行人是否感到舒适、开放、可停留。只有当这些条件相互配合时，街道才更可能从单纯通过空间转化为停留空间。")
    add_para(doc, "因此，在解释结果时，应避免将某一变量简单理解为“越多越好”。例如，POI 总量可能在站前高密度地区与停留高度相关，但在远站区域，设施质量、类型和与步行路径的关系可能比数量更重要。街景绿化也不一定单独决定停留，只有与可达性、界面活跃度和停留目的结合，才可能形成稳定活动。")
    add_heading(doc, "6.3 与金森研究的关系", 2)
    add_para(doc, "金森研究已经证明，距站距离、步道结构和沿道设施会影响步行者停留，并指出远离车站的区域也可能因良好的物理条件和设施条件而产生停留。本研究并不否定这一结论，而是在其基础上进一步扩展变量体系和解释层次。")
    add_para(doc, "具体而言，本文的扩展主要体现在两个方面。第一，引入空间句法指标，将街道在整体网络中的拓扑位置纳入解释，补充金森研究中对路网结构的不足。第二，引入街景语义指标，将行人眼平视角的微观环境纳入模型，补充仅使用步道结构和 POI 难以捕捉的体验性因素。由此，本研究试图从“空间结构-功能配置-微观体验”三个层次解释 TOD 站域停留行为。")
    add_heading(doc, "6.4 规划启示", 2)
    add_para(doc, "若研究结果显示远站区域中某些环境条件能够缓和距离衰减，则大井町站周边更新不应只集中于站前广场和大型开发地块，也应关注连接车站、商店街、住宅区和公共设施的街道网络。对于远离车站但具有较好网络位置的街道，可通过提升沿道设施质量、改善步行空间连续性、增加绿化和休憩支持设施，扩大站域公共生活的空间范围。")
    add_para(doc, "从 TOD 规划角度看，站域活力不应只被理解为高密度开发和高人流集聚，而应被理解为人在站域中愿意停留、使用和互动的能力。通过识别停留行为的空间机制，可以为街道更新、设施诱导和步行环境改善提供更精细的空间依据。")

    add_heading(doc, "第7章 结论", 1)
    add_heading(doc, "7.1 主要结论", 2)
    add_para(doc, "本文以大井町 TOD 站域为对象，将 GPS 停留点作为站域 Place 功能的行为代理指标，并在 street segment 尺度上整合距站距离、空间句法、POI 和街景环境指标，构建了分析停留行为空间分布及建成环境影响机制的研究框架。")
    add_para(doc, "待模型完成后，本节应用 3-4 点概括实证发现：第一，停留行为是否呈现明显站前集中和距离衰减；第二，哪些空间句法、POI 或街景变量显著影响停留；第三，哪些环境条件能够缓和远站距离衰减；第四，本研究相较于金森研究的新增解释价值。")
    add_heading(doc, "7.2 研究不足与未来展望", 2)
    add_para(doc, "本研究仍可能存在若干限制。首先，若仅以大井町为单案例，结果的外推性有限，未来可加入吉祥寺或其他站域进行比较。其次，GPS 数据、POI 数据和街景数据可能存在时间不一致问题。第三，停留点识别受到采样间隔、定位精度和阈值设定影响，需要通过敏感性分析提高可靠性。第四，若未纳入车站出入口、天气、活动事件和地形等因素，模型仍可能遗漏部分解释变量。")
    add_para(doc, "未来研究可以进一步比较不同站域类型下停留机制的差异，也可以结合实地观察、问卷或访谈，解释 GPS 停留点背后的具体活动类型和使用者感受。")

    add_heading(doc, "参考文献（初稿）", 1)
    refs = [
        "Bertolini, L. (1999). Spatial Development Patterns and Public Transport: The Application of an Analytical Model in the Netherlands. Planning Practice and Research, 14(2), 199-210.",
        "Chorus, P., & Bertolini, L. (2011). An application of the node-place model to explore the spatial development dynamics of station areas in Tokyo. Journal of Transport and Land Use, 4(1), 45-58.",
        "Vale, D. S. (2015). Transit-oriented development, integration of land use and transport, and pedestrian accessibility. Journal of Transport Geography, 45, 70-80.",
        "Jones, P., & Boujenko, N. (2009). Link and Place: A new approach to street planning and design. Australasian Transport Research Forum.",
        "Hillier, B., Penn, A., Hanson, J., Grajewski, T., & Xu, J. (1993). Natural Movement: Or, Configuration and Attraction in Urban Pedestrian Movement. Environment and Planning B, 20(1), 29-66.",
        "Zheng, Y., Zhang, L., Xie, X., & Ma, W.-Y. (2009). Mining Interesting Locations and Travel Sequences from GPS Trajectories. WWW 2009, 791-800.",
        "Nagata, S., Nakaya, T., Hanibuchi, T., Amagasa, S., Kikuchi, H., & Inoue, S. (2020). Objective scoring of streetscape walkability related to leisure walking. Health & Place, 66, 102428.",
        "Ki, D., & Lee, S. (2021). Analyzing the effects of Green View Index of neighborhood streets on walking time using Google Street View and deep learning. Landscape and Urban Planning, 205, 103920.",
        "金森貴洋・中山俊・厳網林. 駅からの距離を考慮した歩道構造と沿道施設が滞在数に及ぼす影響：大井町駅と吉祥寺駅周辺部を対象として.",
        "国土交通省（2024）歩行空間ネットワークデータ整備仕様.",
    ]
    for ref in refs:
        add_para(doc, ref)

    doc.save(OUT)


if __name__ == "__main__":
    build()
